"""Inventory a fillable PDF and a spreadsheet, to write a mapping for a new form.

Prints every PDF field (its label on the page, type, hidden/disguised flags,
answer options, the fields its script reads, and the script itself) and every
spreadsheet input and formula (with the text on its row).

Usage: uv run python inspect_sources.py FORM.pdf BOOK.xlsx [SHEET]
"""

import sys

import pymupdf
from openpyxl import load_workbook

from reconcile import extract_pdf, extract_xlsx


def page_labels(path):
    """For each widget, the nearest text to its left on the same line."""
    labels = {}
    for page in pymupdf.open(path):
        words = page.get_text("words")  # x0, y0, x1, y1, word, block, line, n
        for w in page.widgets():
            r = w.rect
            row = [t for t in words if t[3] > r.y0 - 4 and t[1] < r.y1 + 4 and t[2] <= r.x0 + 1]
            row.sort(key=lambda t: (t[5], t[6], t[0]))
            text = " ".join(t[4] for t in row).strip()
            labels[w.field_name] = text[-90:]
    return labels


def main(pdf_path, xlsx_path, sheet=None):
    form = extract_pdf(pdf_path)
    labels = page_labels(pdf_path)
    print(f"# PDF: {pdf_path}")
    print(f"{len(form['fields'])} fields; calculation order: {', '.join(form['calculation_order'])}\n")
    for js in form["document_js"]:
        print("Document-level script:\n" + js)
    for name, f in form["fields"].items():
        flags = [k for k in ("hidden", "white_text", "in_margin") if f.get(k)]
        print(f"## {name}  ({f['type']}{', ' + ', '.join(flags) if flags else ''})")
        print(f"   label on page: {labels.get(name, '')!r}")
        if "options" in f:
            print("   options (value: shown): " + "; ".join(f"{e}: {d}" for e, d in f["options"]))
        if f["calculate"]:
            print(f"   reads: {', '.join(f['depends_on'])}")
            print("   script: " + f["calculate"].replace("\n", "\n           "))
    print()

    wb = load_workbook(xlsx_path)
    sheets = [sheet] if sheet else wb.sheetnames
    for name in sheets:
        ws = wb[name]
        data = extract_xlsx(xlsx_path, name)
        print(f"# Spreadsheet: {xlsx_path}, sheet {name!r}")
        for ref, opts in data["choices"].items():
            print(f"   input {ref}: one of {opts}")
        for ref, formula in data["formulas"].items():
            row = ws[ref].row
            text = " | ".join(str(c.value) for c in ws[row] if isinstance(c.value, str) and not c.value.startswith("="))
            print(f"   {ref}: {formula}    [row: {text[:80]}]")
        print()


if __name__ == "__main__":
    main(*sys.argv[1:4])
