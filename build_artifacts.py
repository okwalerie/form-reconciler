"""Build the two synthetic artifacts: a legacy fillable PDF and an Excel calculator.

The PDF imitates a hand-built Acrobat form: calculations live in per-field
calculate actions, a helper function sits in document-level JavaScript, the
scores are hidden fields (the form only shows levels), and a white-text
helper field in the left margin holds an intermediate value.
"""

import pikepdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from procedure import BANDS, CATEGORIES, CATEGORY_FR, FACTORS, UNC_LEVELS, UNCERTAINTY, WEIGHTS

PDF = "artifacts/field-visit-risk-assessment.pdf"
XLSX = "artifacts/risk-factor-values.xlsx"

# The PDF grades its (multiplied) uncertainty score on its own scale.
PDF_UNC_LEVELS = [(6_000, "Low"), (21_000, "Medium"), (float("inf"), "High")]


# --- PDF -------------------------------------------------------------------

DOC_JS = """\
// Shared helpers. Added 2016, see calc order if totals look wrong.
var UNC = {%s};
function sumFields(doc, names) {
  var s = 0;
  for (var i = 0; i < names.length; i++) s += Number(doc.getField(names[i]).value);
  return s;
}
""" % ", ".join(f'"{k}": {v}' for k, v in UNCERTAINTY.items())


def level_script(source, levels):
    lines = [f'var t = Number(this.getField("{source}").value);']
    for i, (limit, name) in enumerate(levels[:-1]):
        lines.append(f'{"else " if i else ""}if (t < {limit}) event.value = "{name}";')
    lines.append(f'else event.value = "{levels[-1][1]}";')
    return "\n".join(lines)


def calc_scripts():
    """Field name -> calculate script, in calculation order."""
    js = {}
    for fid, _, _, _, wcat, _ in FACTORS:
        w = WEIGHTS[wcat]
        js[f"risk_{fid}"] = f'event.value = Number(this.getField("opt_{fid}").value) * {w};'
        js[f"unc_{fid}"] = f'event.value = UNC[this.getField("opt_{fid}").value] * {w};'

    # Hidden margin helper. Recomputes the likelihood subtotal with the
    # weights it was written against, and the total reads it.
    js["zz_L"] = (
        "// L subtotal (weights per 2016 table)\n"
        "var s = 0;\n"
        's += this.getField("opt_L1").value * 1;\n'
        's += this.getField("opt_L2").value * 3;\n'
        's += this.getField("opt_L3").value * 5;\n'
        "event.value = s;"
    )
    for cat in CATEGORIES:
        ids = [f[0] for f in FACTORS if f[1] == cat]
        names = ", ".join(f'"risk_{i}"' for i in ids)
        js[f"sub_{cat[0]}"] = f"event.value = sumFields(this, [{names}]);"
    js["total_risk"] = (
        'event.value = this.getField("zz_L").value * this.getField("sub_C").value'
        ' * this.getField("sub_M").value;'
    )
    js["band"] = level_script("total_risk", BANDS)
    unc_parts = []
    for cat in CATEGORIES:
        ids = [f[0] for f in FACTORS if f[1] == cat]
        unc_parts.append("sumFields(this, [" + ", ".join(f'"unc_{i}"' for i in ids) + "])")
    js["unc_total"] = "event.value = " + " * ".join(unc_parts) + ";"
    js["unc_level"] = level_script("unc_total", PDF_UNC_LEVELS)
    js["warning"] = (
        'var w = this.getField("opt_L1").value == "10";\n'
        "// Region B: hazardous materials warning from 'Possible' up\n"
        'if (Number(this.getField("opt_C1").value) >= 3) w = true;\n'
        'event.value = w ? "WARNING / AVERTISSEMENT" : "";'
    )
    js["approval"] = (
        'var b = this.getField("band").value;\n'
        'var w = this.getField("warning").value != "";\n'
        'event.value = (w || b == "High" || b == "Extreme") ? "Yes / Oui" : "No / Non";'
    )
    return js


def build_pdf():
    c = canvas.Canvas(PDF, pagesize=letter)
    form = c.acroForm
    width, height = letter
    c.setTitle("Field Visit Risk Assessment / Évaluation des risques de visite sur le terrain")

    c.setFont("Helvetica-Bold", 13)
    c.drawString(54, height - 50, "Field Visit Risk Assessment")
    c.setFont("Helvetica", 10)
    c.drawString(54, height - 64, "Évaluation des risques de visite sur le terrain")
    c.setFont("Helvetica", 7)
    c.drawRightString(width - 54, height - 50, "Form FVRA-04 (rev. 2016-03)")
    c.setFont("Helvetica-Oblique", 7)
    c.drawString(54, height - 76, "SYNTHETIC FORM FOR DEMONSTRATION. Fictional program; not a real government form.")

    c.setFont("Helvetica-Bold", 7)
    y = height - 100
    for x, label in [(54, "Factor / Facteur"), (330, "Weight / Poids"), (400, "Answer / Réponse")]:
        c.drawString(x, y, label)

    def hidden_number(name, x, y):
        # Scores exist on the form but are never shown.
        form.textfield(name=name, x=x, y=y, width=30, height=12, fontSize=7,
                       fieldFlags="readOnly", borderWidth=0, annotationFlags="hidden")

    y -= 8
    for cat in CATEGORIES:
        y -= 18
        c.setFont("Helvetica-Bold", 8)
        c.drawString(54, y, f"{cat} / {CATEGORY_FR[cat]}")
        for fid, fcat, en, fr, wcat, options in FACTORS:
            if fcat != cat:
                continue
            y -= 26
            c.setFont("Helvetica", 8)
            c.drawString(54, y + 8, en)
            c.setFont("Helvetica", 7)
            c.drawString(54, y - 1, fr)
            c.setFont("Helvetica", 8)
            c.drawString(330, y + 4, f"{wcat} ({WEIGHTS[wcat]})")
            form.choice(
                name=f"opt_{fid}", x=380, y=y - 2, width=130, height=16, fontSize=7,
                options=[(f"{o_en} / {o_fr}", str(s)) for o_en, o_fr, s in options],
                value=str(options[0][2]),
            )
            hidden_number(f"risk_{fid}", 520, y)
            hidden_number(f"unc_{fid}", 520, y - 12)
        hidden_number(f"sub_{cat[0]}", 520, y - 26)
        y -= 8

    y -= 34
    for name, label in [
        ("band", "Overall risk / Risque global"),
        ("unc_level", "Uncertainty level / Niveau d'incertitude"),
        ("warning", "Single critical factor / Facteur critique unique"),
        ("approval", "Supervisor approval required / Approbation requise"),
    ]:
        c.setFont("Helvetica-Bold", 8)
        c.drawString(54, y + 4, label)
        form.textfield(name=name, x=330, y=y - 2, width=180, height=16, fontSize=8,
                       fieldFlags="readOnly", borderWidth=0, fillColor=colors.whitesmoke)
        y -= 22
    hidden_number("total_risk", 520, y + 22)
    hidden_number("unc_total", 520, y + 44)

    # The margin helper: white on white, no border, outside the printed table.
    form.textfield(name="zz_L", x=6, y=height - 330, width=30, height=10, fontSize=6,
                   fieldFlags="readOnly", borderWidth=0, fillColor=colors.white,
                   textColor=colors.white, borderColor=colors.white)
    c.save()

    scripts = calc_scripts()
    with pikepdf.open(PDF, allow_overwriting_input=True) as pdf:
        acro = pdf.Root.AcroForm
        by_name = {str(f.T): f for f in acro.Fields}
        for name, js in scripts.items():
            f = by_name[name]
            f.AA = pikepdf.Dictionary(C=pikepdf.Dictionary(S=pikepdf.Name.JavaScript, JS=pikepdf.String(js)))
        acro.CO = pikepdf.Array([by_name[n] for n in scripts])
        acro.NeedAppearances = True
        pdf.Root.Names = pikepdf.Dictionary(
            JavaScript=pikepdf.Dictionary(Names=pikepdf.Array([
                pikepdf.String("helpers"),
                pdf.make_indirect(pikepdf.Dictionary(S=pikepdf.Name.JavaScript, JS=pikepdf.String(DOC_JS))),
            ]))
        )
        pdf.save(PDF)


# --- Spreadsheet -----------------------------------------------------------

def nested_if(cell, levels):
    expr = f'"{levels[-1][1]}"'
    for limit, name in reversed(levels[:-1]):
        expr = f'IF({cell}<{limit},"{name}",{expr})'
    return "=" + expr


def build_xlsx():
    wb = Workbook()
    bold = Font(bold=True)
    inp = PatternFill("solid", fgColor="FFF2CC")

    ws_w = wb.active
    ws_w.title = "Weights"
    ws_w.append(["Weight category", "Weight"])
    for k, v in WEIGHTS.items():
        ws_w.append([k, v])

    ws_o = wb.create_sheet("Options")
    ws_o.append(["Key", "Factor", "Option", "Risk score", "Uncertainty score"])
    for fid, _, _, _, _, options in FACTORS:
        for o_en, _, s in options:
            ws_o.append([f"{fid}|{o_en}", fid, o_en, s, UNCERTAINTY[s]])
    n_opt = ws_o.max_row

    ws = wb.create_sheet("Assessment", 0)
    ws["A1"] = "Field Visit Risk Assessment: risk factor values (SYNTHETIC, fictional program)"
    ws["A1"].font = bold
    ws.append([])
    ws.append(["ID", "Category", "Factor", "Weight category", "Answer", "Weight", "Risk score",
               "Uncertainty", "Row risk", "Row uncertainty", "Critical warning"])
    for cell in ws[3]:
        cell.font = bold
    first = 4
    for i, (fid, cat, en, _, wcat, options) in enumerate(FACTORS):
        r = first + i
        ws.append([
            fid, cat, en, wcat, options[0][0],
            f"=VLOOKUP(D{r},Weights!$A$2:$B$5,2,FALSE)",
            f'=INDEX(Options!$D$2:$D${n_opt},MATCH(A{r}&"|"&E{r},Options!$A$2:$A${n_opt},0))',
            f'=INDEX(Options!$E$2:$E${n_opt},MATCH(A{r}&"|"&E{r},Options!$A$2:$A${n_opt},0))',
            f"=G{r}*F{r}",
            f"=H{r}*F{r}",
            f'=IF(AND(D{r}="Critical",G{r}=10),1,0)',
        ])
        ws[f"E{r}"].fill = inp
        dv = DataValidation(type="list", formula1='"' + ",".join(o[0] for o in options) + '"')
        ws.add_data_validation(dv)
        dv.add(f"E{r}")
    last = first + len(FACTORS) - 1

    r = last + 2
    rows = {}
    for cat in CATEGORIES:
        ws[f"H{r}"] = f"{cat} subtotal"
        ws[f"I{r}"] = f'=MAX(1,SUMIF($B${first}:$B${last},"{cat}",$I${first}:$I${last}))'
        rows[cat] = r
        r += 1
    r += 1
    out = {}
    for label in ["Total risk score", "Overall risk", "Uncertainty score", "Uncertainty level",
                  "Single critical factor warning", "Supervisor approval required"]:
        ws[f"H{r}"] = label
        ws[f"H{r}"].font = bold
        out[label] = f"I{r}"
        r += 1
    ws[out["Total risk score"]] = "=" + "*".join(f"I{rows[c]}" for c in CATEGORIES)
    ws[out["Overall risk"]] = nested_if(out["Total risk score"], BANDS)
    ws[out["Uncertainty score"]] = f"=SUM(J{first}:J{last})"
    ws[out["Uncertainty level"]] = nested_if(out["Uncertainty score"], UNC_LEVELS)
    ws[out["Single critical factor warning"]] = f'=IF(SUM(K{first}:K{last})>0,"WARNING","")'
    b, w = out["Overall risk"], out["Single critical factor warning"]
    ws[out["Supervisor approval required"]] = f'=IF(OR({w}="WARNING",{b}="High",{b}="Extreme"),"Yes","No")'
    for col, wd in zip("ABCDEFGHIJK", [5, 12, 30, 15, 20, 8, 10, 11, 26, 14, 14]):
        ws.column_dimensions[col].width = wd
    wb.save(XLSX)
    return out


if __name__ == "__main__":
    import os
    os.makedirs("artifacts", exist_ok=True)
    build_pdf()
    print("built", PDF)
    print("built", XLSX, build_xlsx())
